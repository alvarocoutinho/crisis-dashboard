"""Monta a amostra de anotação do estudo de calibração dos filtros (v0.3.3 → v0.4).

Entrada: os arquivos do notebook notebooks/ODRD_coleta_retrospectiva.ipynb
(Drive do pesquisador; títulos e descrições não vão para o repositório):
  resultados_alertas.csv, videos.csv, [videos_complemento.csv], [pasta de transcrições]

Unidade de anotação: o PAR alerta–vídeo (o mesmo vídeo pode ser municipal para
um alerta e regional para outro). Todos os candidatos de todas as consultas
(padrão e variantes) formam o conjunto julgado (pooling).

Estratos (definidos pela decisão ATUAL dos filtros, que é o que se quer medir):
  E1 aceito, cita o município            E2 aceito, não cita (regional)
  E3 rejeitado, município no título      E4 rejeitado, município só na descrição
  E5 rejeitado, sem menção, termo de perigo no título
  E6 rejeitado, sem menção, sem termo de perigo (teste de elusão)
  E7 só recuperado por consulta variante
Cada par sorteado leva peso = tamanho do estrato / sorteados no estrato, para
estimar precisão e recuperação da população (não só da amostra).

Saída (pasta --saida, fica no Drive):
  anotacao_codificador1.xlsx  — planilha cega (sem decisão nem score), ordem aleatória
  anotacao_codificador2.xlsx  — subconjunto para dupla codificação (confiabilidade)
  anotacao_treino.xlsx        — 30 pares de prática, fora da amostra
  anotacao_validacao_assistida.xlsx — 50 pares da amostra para o codificador 1
                                re-rotular assistindo ao vídeo (erro do julgamento por metadados)
  anotacao_chave.csv          — id_par → estrato, peso, decisão e scores (só o pesquisador)
  ia/                         — prompts prontos por lote para os anotadores de IA
                                (condições metadados e transcricao) + manifesto com hashes

Uso:
  python scripts/build_annotation_sample.py --dados PASTA_DO_DRIVE --saida PASTA \\
      [--transcricoes PASTA/texto] [--n 700] [--dupla 150] [--semente 20261014]
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

REPO = Path(__file__).parent.parent.resolve()
sys.path.insert(0, str(REPO))

from faa_poc.constants import EVENT_TERMS  # noqa: E402
from faa_poc.text import normalize  # noqa: E402

COTAS = {"E1": 150, "E2": 120, "E3": 150, "E4": 80, "E5": 80, "E6": 60, "E7": 60}
HAZARD = sorted({normalize(t) for fam, ts in EVENT_TERMS.items() if fam != "default" for t in ts}, key=len, reverse=True)
ROTULOS = {
    "perigo": ["sim", "nao", "incerto"],
    "episodio": ["sim", "nao", "incerto"],
    "lugar": ["municipio", "regiao_vizinho", "estado_outro", "indeterminado"],
    "fonte_decisao": ["metadados", "transcricao", "assistiu"],
}


def _ler(p: Path) -> list[dict]:
    if not p.exists():
        return []
    with p.open(encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def _hazard(text: str) -> bool:
    t = normalize(text)
    return any(re.search(rf"(?<![a-z]){re.escape(h)}(?![a-z])", t) for h in HAZARD)


def _kwic(texto: str, termos: list[str], janela: int = 25, maximo: int = 3) -> str:
    """Trechos da transcrição em volta do município e dos termos de perigo."""
    palavras = texto.split()
    norm = [normalize(w) for w in palavras]
    alvos = [normalize(t).split() for t in termos if t]
    trechos, ult = [], -10**9
    for i in range(len(norm)):
        if any(norm[i:i + len(a)] == a for a in alvos) and i - ult > janela:
            trechos.append("… " + " ".join(palavras[max(0, i - janela): i + janela]) + " …")
            ult = i
            if len(trechos) >= maximo:
                break
    return "\n".join(trechos)


def montar_pool(dados: Path) -> list[dict]:
    res = _ler(dados / "resultados_alertas.csv")
    ultimo: dict[str, dict] = {}
    for r in res:
        if r.get("status") == "ok" and r.get("coletado_em", "") >= ultimo.get(r["alert_id"], {}).get("coletado_em", ""):
            ultimo[r["alert_id"]] = r
    compl = {r["video_id"]: r for r in _ler(dados / "videos_complemento.csv")}
    vids = _ler(dados / "videos.csv")
    # lote 1 não tinha coletado_em nos vídeos: vale a linha sem data se o resultado mais recente for do lote 1
    tem_data = {(r["alert_id"], r.get("coletado_em", "")) for r in vids if r.get("coletado_em")}
    pares: dict[tuple[str, str], dict] = {}
    for v in vids:
        a = ultimo.get(v["alert_id"])
        if not a:
            continue
        var = v.get("variante") or "padrao"
        if var == "padrao":
            ok = (v.get("coletado_em") == a["coletado_em"]) or (not v.get("coletado_em") and (v["alert_id"], a["coletado_em"]) not in tem_data)
            if not ok:
                continue
        c = compl.get(v["video_id"], {})
        chave = (v["alert_id"], v["video_id"])
        p = pares.setdefault(chave, {
            "alert_id": v["alert_id"], "video_id": v["video_id"], "municipio": a["municipio"], "uf": a["uf"],
            "evento": a.get("evento", ""), "alerta_aberto_em": a.get("opened_at", ""), "grupo": a.get("grupo", ""),
            "title": v.get("title", ""), "channel_title": v.get("channel_title", ""),
            "published_at": v.get("published_at") or c.get("published_at", ""),
            "description": v.get("description") or c.get("description", "") or v.get("description_snippet", ""),
            "municipal_evidence": v.get("municipal_evidence", ""), "aceito_padrao": None, "variantes": set(),
            "relevance_score": v.get("relevance_score", ""), "geo_score": v.get("geo_score", ""),
            "event_score": v.get("event_score", ""), "rejection_reason": v.get("rejection_reason", ""),
        })
        p["variantes"].add(var)
        if var == "padrao":
            p["aceito_padrao"] = v.get("aceito") == "1"
            p.update({k: v.get(k, "") for k in ("relevance_score", "geo_score", "event_score", "rejection_reason")})
    for p in pares.values():
        ev = p["municipal_evidence"]
        if p["aceito_padrao"] is None:
            p["estrato"] = "E7"
        elif p["aceito_padrao"]:
            p["estrato"] = "E1" if ev not in ("", "nenhuma") else "E2"
        elif "titulo" in ev:
            p["estrato"] = "E3"
        elif ev == "descricao":
            p["estrato"] = "E4"
        else:
            p["estrato"] = "E5" if _hazard(p["title"]) else "E6"
        p["variantes"] = ";".join(sorted(p["variantes"]))
    return list(pares.values())


def sortear(pool: list[dict], n: int, rng: random.Random) -> list[dict]:
    por = defaultdict(list)
    for p in pool:
        por[p["estrato"]].append(p)
    escala = n / sum(COTAS.values())
    cotas = {e: min(len(por[e]), round(q * escala)) for e, q in COTAS.items()}
    sobra = n - sum(cotas.values())
    while sobra > 0:  # redistribui para estratos que ainda têm pares
        abertos = [e for e in COTAS if len(por[e]) > cotas[e]]
        if not abertos:
            break
        for e in abertos:
            if sobra == 0:
                break
            cotas[e] += 1; sobra -= 1
    amostra = []
    for e, q in cotas.items():
        for p in rng.sample(por[e], q):
            amostra.append({**p, "peso": round(len(por[e]) / q, 4)})
    return amostra


def _xlsx(caminho: Path, linhas: list[dict], transcricoes: Path | None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); ws = wb.active; ws.title = "anotacao"
    cols = ["id_par", "municipio", "uf", "evento", "alerta_aberto_em", "title", "channel_title", "published_at",
            "description", "transcricao_trechos", "url", *ROTULOS, "observacao"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in linhas:
        trechos = ""
        if transcricoes is not None and (transcricoes / f"{r['video_id']}.txt").exists():
            texto = (transcricoes / f"{r['video_id']}.txt").read_text(encoding="utf-8", errors="ignore")
            trechos = _kwic(texto, [r["municipio"], *[h for h in HAZARD if " " not in h][:40]]) or "(transcrição sem o município nem termos de perigo)"
        elif transcricoes is not None:
            trechos = "(sem transcrição)"
        ws.append([r["id_par"], r["municipio"], r["uf"], r["evento"], r["alerta_aberto_em"][:16], r["title"],
                   r["channel_title"], r["published_at"][:16], (r["description"] or "")[:1500], trechos,
                   f"https://www.youtube.com/watch?v={r['video_id']}", "", "", "", "", ""])
    n = len(linhas) + 1
    for i, (campo, opcoes) in enumerate(ROTULOS.items()):
        col = chr(ord("A") + cols.index(campo))
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n}")
    for col, larg in zip("ABCDEFGHIJK", (9, 18, 5, 18, 16, 60, 22, 16, 70, 60, 30)):
        ws.column_dimensions[col].width = larg
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ref = wb.create_sheet("rotulos")
    ref.append(["campo", "valores (ver docs/calibracao/livro_de_codigos.md)"])
    for k, v in ROTULOS.items():
        ref.append([k, " | ".join(v)])
    wb.save(caminho)


TRANSCRICAO_MAX = 12000


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def exportar_ia(saida: Path, grupos: dict[str, list[dict]], transcricoes: Path | None,
                tam_lote: dict[str, int]) -> None:
    """Prompts prontos por lote para os anotadores de IA (prompt congelado v1).

    metadados: o mesmo material que os humanos veem (sem trechos de transcrição).
    transcricao: só a fala transcrita, sem título nem descrição; pares sem
    transcrição ficam de fora (dado faltante, aceito pelo alfa).
    """
    prompt_p = REPO / "docs/calibracao/prompt_ia_v1.md"
    livro_p = REPO / "docs/calibracao/livro_de_codigos.md"
    modelo = prompt_p.read_text(encoding="utf-8").split("## Texto enviado", 1)[1].strip()
    livro = livro_p.read_text(encoding="utf-8")
    pasta = saida / "ia"; (pasta / "prompts").mkdir(parents=True, exist_ok=True)
    manifesto = {"prompt": "prompt_ia_v1.md", "sha256_prompt": _sha(prompt_p), "sha256_livro": _sha(livro_p),
                 "transcricao_max_caracteres": TRANSCRICAO_MAX, "lotes": []}
    for grupo, linhas in grupos.items():
        for condicao in ("metadados", "transcricao"):
            itens = []
            for r in linhas:
                base = {"id_par": r["id_par"], "municipio": r["municipio"], "uf": r["uf"], "evento": r["evento"],
                        "alerta_aberto_em": r["alerta_aberto_em"][:16]}
                if condicao == "metadados":
                    itens.append({**base, "titulo": r["title"], "canal": r["channel_title"],
                                  "publicado_em": r["published_at"][:16], "descricao": (r["description"] or "")[:1500]})
                elif transcricoes is not None and (transcricoes / f"{r['video_id']}.txt").exists():
                    texto = (transcricoes / f"{r['video_id']}.txt").read_text(encoding="utf-8", errors="ignore")
                    itens.append({**base, "transcricao": texto[:TRANSCRICAO_MAX]})
            tam = tam_lote[condicao]  # transcrições são longas: lotes menores
            for i in range(0, len(itens), tam):
                lote = itens[i:i + tam]
                nome = f"{grupo}_{condicao}_lote{i // tam + 1:02d}"
                texto = (modelo.replace("{LIVRO_DE_CODIGOS}", livro).replace("{CONDICAO}", condicao)
                         .replace("{N}", str(len(lote)))
                         .replace("{LOTE_JSONL}", "\n".join(json.dumps(x, ensure_ascii=False) for x in lote)))
                (pasta / "prompts" / f"{nome}.txt").write_text(texto, encoding="utf-8")
                manifesto["lotes"].append({"lote": nome, "grupo": grupo, "condicao": condicao,
                                           "ids": [x["id_par"] for x in lote]})
    (pasta / "manifesto.json").write_text(json.dumps(manifesto, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dados", type=Path, required=True)
    ap.add_argument("--saida", type=Path, required=True)
    ap.add_argument("--transcricoes", type=Path)
    ap.add_argument("--n", type=int, default=700)
    ap.add_argument("--dupla", type=int, default=150)
    ap.add_argument("--treino", type=int, default=30)
    ap.add_argument("--assistidos", type=int, default=50)
    # piloto de 14/10: um subagente com 30 itens de metadados custou ~85 mil tokens,
    # quase todo fixo (instruções + livro); lotes maiores diluem esse custo
    ap.add_argument("--lote-ia-metadados", type=int, default=60)
    ap.add_argument("--lote-ia-transcricao", type=int, default=8)
    ap.add_argument("--semente", type=int, default=20261014)
    args = ap.parse_args()
    rng = random.Random(args.semente)

    pool = montar_pool(args.dados)
    tam = Counter(p["estrato"] for p in pool)
    print(f"pool: {len(pool)} pares alerta–vídeo; por estrato: {dict(sorted(tam.items()))}")
    amostra = sortear(pool, args.n, rng)
    rng.shuffle(amostra)  # ordem aleatória: o codificador não vê os estratos
    for i, p in enumerate(amostra, 1):
        p["id_par"] = f"P{i:04d}"
    escolhidos = {(p["alert_id"], p["video_id"]) for p in amostra}
    resto = [p for p in pool if (p["alert_id"], p["video_id"]) not in escolhidos]
    treino = rng.sample(resto, min(args.treino, len(resto)))
    for i, p in enumerate(treino, 1):
        p["id_par"] = f"T{i:03d}"
    dupla = rng.sample(amostra, min(args.dupla, len(amostra)))  # aleatória simples: alfa sobre a distribuição real
    dupla.sort(key=lambda p: p["id_par"])
    assistidos = sorted(rng.sample(amostra, min(args.assistidos, len(amostra))), key=lambda p: p["id_par"])

    args.saida.mkdir(parents=True, exist_ok=True)
    _xlsx(args.saida / "anotacao_codificador1.xlsx", amostra, args.transcricoes)
    _xlsx(args.saida / "anotacao_codificador2.xlsx", dupla, args.transcricoes)
    _xlsx(args.saida / "anotacao_treino.xlsx", treino, args.transcricoes)
    _xlsx(args.saida / "anotacao_validacao_assistida.xlsx", assistidos, args.transcricoes)
    exportar_ia(args.saida, {"treino": treino, "amostra": sorted(amostra, key=lambda p: p["id_par"])},
                args.transcricoes, {"metadados": args.lote_ia_metadados, "transcricao": args.lote_ia_transcricao})
    campos = ["id_par", "alert_id", "video_id", "estrato", "peso", "grupo", "municipio", "uf", "evento",
              "aceito_padrao", "variantes", "municipal_evidence", "relevance_score", "geo_score", "event_score",
              "rejection_reason", "dupla", "assistido", "tem_transcricao"]
    ids_dupla = {p["id_par"] for p in dupla}
    ids_assist = {p["id_par"] for p in assistidos}
    with (args.saida / "anotacao_chave.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        for p in sorted(amostra + treino, key=lambda p: p["id_par"]):
            w.writerow({**p, "dupla": int(p["id_par"] in ids_dupla), "assistido": int(p["id_par"] in ids_assist),
                        "tem_transcricao": int(args.transcricoes is not None
                                               and (args.transcricoes / f"{p['video_id']}.txt").exists()),
                        "peso": p.get("peso", "")})
    sorteados = Counter(p["estrato"] for p in amostra)
    print(f"amostra: {len(amostra)} pares {dict(sorted(sorteados.items()))}; dupla: {len(dupla)}; "
          f"treino: {len(treino)}; validação assistida: {len(assistidos)}")
    print(f"gravado em {args.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""Análise do estudo de calibração: confiabilidade, gabarito e desempenho dos filtros e das IAs.

Entradas (pasta --dir, gerada por build_annotation_sample.py e preenchida):
  anotacao_chave.csv                      (do amostrador)
  anotacao_codificador1.xlsx              (preenchida pelo codificador 1)
  anotacao_codificador2.xlsx              (preenchida pelo codificador 2, pares em dupla)
  anotacao_validacao_assistida.xlsx       (opcional: codificador 1 assistindo aos vídeos)
  adjudicacao.csv                         (opcional: id_par, dimensao, valor — resolve divergências da dupla)
  ia/resultados/rodada{1,2}/*.jsonl       (opcional: saídas dos anotadores de IA, uma linha JSON por par)

Saída: relatorio_calibracao.md + metricas.csv na mesma pasta.

Regras:
  - O gabarito é humano: codificador 1, com as divergências da dupla substituídas
    pela adjudicação (pares em dupla sem adjudicação ficam fora do gabarito).
  - A IA nunca entra no gabarito; é avaliada contra ele.
  - Precisão e recuperação dos filtros usam os pesos dos estratos (estimativa
    para a população de candidatos, não só para a amostra).

Uso: python scripts/analyze_annotation.py --dir PASTA [--bootstrap 1000]
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import random
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path

DIMENSOES = ("perigo", "episodio", "lugar")


# ---------------------------------------------------------------- alfa de Krippendorff
def alfa_nominal(unidades: list[list[str]]) -> float:
    """Alfa de Krippendorff, métrica nominal, com dados faltantes.

    `unidades`: para cada item, a lista de valores atribuídos (só os presentes).
    Itens com menos de 2 valores não são pareáveis e ficam de fora.
    """
    o = Counter()  # matriz de coincidências o[(c, k)]
    for vals in unidades:
        vals = [v for v in vals if v not in (None, "")]
        m = len(vals)
        if m < 2:
            continue
        for i, j in ((i, j) for i in range(m) for j in range(m) if i != j):
            o[(vals[i], vals[j])] += 1 / (m - 1)
    n_c = Counter()
    for (c, _k), v in o.items():
        n_c[c] += v
    n = sum(n_c.values())
    if n <= 1:
        return float("nan")
    d_o = sum(v for (c, k), v in o.items() if c != k)
    d_e = sum(n_c[c] * n_c[k] for c in n_c for k in n_c if c != k) / (n - 1)
    if d_e == 0:
        return 1.0 if d_o == 0 else float("nan")
    return 1 - d_o / d_e


def alfa_ic(unidades: list[list[str]], b: int, rng: random.Random) -> tuple[float, float, float, int]:
    a = alfa_nominal(unidades)
    paraveis = [u for u in unidades if sum(1 for v in u if v not in (None, "")) >= 2]
    if b <= 0 or len(paraveis) < 5:
        return a, float("nan"), float("nan"), len(paraveis)
    amostras = sorted(alfa_nominal([rng.choice(paraveis) for _ in paraveis]) for _ in range(b))
    amostras = [x for x in amostras if not math.isnan(x)]
    return a, amostras[int(0.025 * len(amostras))], amostras[int(0.975 * len(amostras)) - 1], len(paraveis)


# ---------------------------------------------------------------- leitura
def ler_xlsx(caminho: Path) -> dict[str, dict]:
    if not caminho.exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(caminho, read_only=True).active
    linhas = list(ws.iter_rows(values_only=True))
    cab = [str(c) for c in linhas[0]]
    out = {}
    for row in linhas[1:]:
        r = dict(zip(cab, row))
        if r.get("id_par"):
            out[str(r["id_par"])] = {d: (str(r.get(d)).strip() if r.get(d) not in (None, "") else "") for d in (*DIMENSOES, "fonte_decisao")}
    return out


def ler_ia(pasta: Path) -> dict[str, dict]:
    out = {}
    for f in sorted(pasta.glob("*.jsonl")) if pasta.exists() else []:
        condicao = "transcricao" if "_transcricao_" in f.name else "metadados"
        for linha in f.read_text(encoding="utf-8").splitlines():
            linha = linha.strip()
            if not linha.startswith("{"):
                continue
            try:
                r = json.loads(linha)
            except json.JSONDecodeError:
                continue
            out[(condicao, str(r.get("id_par")))] = {d: str(r.get(d, "")).strip() for d in DIMENSOES}
    return out


def alvo(r: dict | None) -> tuple[int | None, int | None]:
    """(municipal_relevante, regional_relevante) a partir dos três rótulos."""
    if not r or any(r.get(d, "") == "" for d in DIMENSOES):
        return None, None
    ok = r["perigo"] == "sim" and r["episodio"] == "sim"
    return int(ok and r["lugar"] == "municipio"), int(ok and r["lugar"] in ("municipio", "regiao_vizinho"))


def prec_rec(pares: list[tuple[int, int, float]]) -> tuple[float, float, float]:
    """(precisão, recuperação, F1) ponderadas; pares = (predito, verdadeiro, peso)."""
    tp = sum(w for p, t, w in pares if p and t)
    fp = sum(w for p, t, w in pares if p and not t)
    fn = sum(w for p, t, w in pares if not p and t)
    pr = tp / (tp + fp) if tp + fp else float("nan")
    rc = tp / (tp + fn) if tp + fn else float("nan")
    f1 = 2 * pr * rc / (pr + rc) if pr + rc and not math.isnan(pr + rc) else float("nan")
    return pr, rc, f1


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", type=Path, required=True)
    ap.add_argument("--bootstrap", type=int, default=1000)
    ap.add_argument("--semente", type=int, default=7)
    args = ap.parse_args()
    d, rng = args.dir, random.Random(args.semente)

    chave = {r["id_par"]: r for r in csv.DictReader((d / "anotacao_chave.csv").open(encoding="utf-8"))}
    amostra = {k: v for k, v in chave.items() if k.startswith("P")}
    c1 = ler_xlsx(d / "anotacao_codificador1.xlsx")
    c2 = ler_xlsx(d / "anotacao_codificador2.xlsx")
    assist = ler_xlsx(d / "anotacao_validacao_assistida.xlsx")
    adj = defaultdict(dict)
    if (d / "adjudicacao.csv").exists():
        for r in csv.DictReader((d / "adjudicacao.csv").open(encoding="utf-8")):
            adj[r["id_par"]][r["dimensao"]] = r["valor"].strip()
    ia1 = ler_ia(d / "ia/resultados/rodada1")
    ia2 = ler_ia(d / "ia/resultados/rodada2")

    # gabarito humano
    gab, pendentes = {}, 0
    for pid in amostra:
        r = dict(c1.get(pid, {}))
        if not r or any(not r.get(x) for x in DIMENSOES):
            continue
        if pid in c2:
            for dim in DIMENSOES:
                if c2[pid].get(dim) and c2[pid][dim] != r[dim]:
                    r[dim] = adj.get(pid, {}).get(dim, "")
            if any(not r[x] for x in DIMENSOES):
                pendentes += 1
                continue
        gab[pid] = r

    metricas, md = [], ["# Relatório do estudo de calibração", ""]
    md += [f"Pares na amostra: {len(amostra)}; rotulados pelo codificador 1: {len(c1)}; "
           f"em dupla: {len(c2)}; no gabarito: {len(gab)}; divergências sem adjudicação: {pendentes}.", ""]

    # 1. confiabilidade e concordância com o gabarito
    comparacoes = [("H1 × H2 (confiabilidade do livro)", lambda p: c1.get(p), lambda p: c2.get(p), [p for p in c2 if p in c1])]
    for cond in ("metadados", "transcricao"):
        comparacoes += [
            (f"IA-{cond} × gabarito", lambda p, c=cond: ia1.get((c, p)), lambda p: gab.get(p), list(gab)),
            (f"IA-{cond} rodada 1 × rodada 2 (autoconsistência)", lambda p, c=cond: ia1.get((c, p)),
             lambda p, c=cond: ia2.get((c, p)), list(amostra)),
        ]
    comparacoes.append(("IA-metadados × IA-transcrição", lambda p: ia1.get(("metadados", p)),
                        lambda p: ia1.get(("transcricao", p)), list(amostra)))
    comparacoes.append(("Metadados × assistindo (codificador 1)", lambda p: c1.get(p), lambda p: assist.get(p), list(assist)))
    md += ["## Alfa de Krippendorff (nominal), por dimensão", "",
           "| Comparação | Dimensão | α | IC95% (bootstrap) | n pares |", "|---|---|---|---|---|"]
    for nome, fa, fb, ids in comparacoes:
        for dim in DIMENSOES:
            un = [[(fa(p) or {}).get(dim, ""), (fb(p) or {}).get(dim, "")] for p in ids]
            a, lo, hi, n = alfa_ic(un, args.bootstrap, rng)
            if n == 0:
                continue
            md.append(f"| {nome} | {dim} | {a:.3f} | {lo:.3f}–{hi:.3f} | {n} |")
            metricas.append({"metrica": "alfa", "comparacao": nome, "dimensao": dim, "valor": round(a, 4),
                             "ic_inf": round(lo, 4), "ic_sup": round(hi, 4), "n": n})
    md.append("")

    # 2. mudança de rótulo ao assistir
    if assist:
        md += ["## Validação assistida", ""]
        for dim in DIMENSOES:
            ids = [p for p in assist if c1.get(p, {}).get(dim) and assist[p].get(dim)]
            mud = sum(1 for p in ids if c1[p][dim] != assist[p][dim])
            md.append(f"- {dim}: {mud}/{len(ids)} rótulos mudaram ao assistir ({(mud / len(ids) if ids else float('nan')):.0%})")
        md.append("")

    # 3. desempenho do filtro atual e das IAs contra o gabarito
    md += ["## Desempenho contra o gabarito (ponderado pelos estratos)", "",
           "| Classificador | Alvo | Precisão | Recuperação | F1 | n |", "|---|---|---|---|---|---|"]
    for i, alvo_nome in enumerate(("municipal_relevante", "regional_relevante")):
        def filtro(p):
            k = chave[p]
            aceito = k.get("aceito_padrao") == "True"
            mun = k.get("municipal_evidence") not in ("", "nenhuma")
            return int(aceito and mun) if i == 0 else int(aceito)
        fontes = {"filtro atual (v0.3.3)": filtro}
        for cond in ("metadados", "transcricao"):
            fontes[f"IA-{cond}"] = lambda p, c=cond: alvo(ia1.get((c, p)))[i]
        for nome, f in fontes.items():
            pares = [(f(p), alvo(gab[p])[i], float(chave[p]["peso"] or 1)) for p in gab]
            pares = [x for x in pares if x[0] is not None and x[1] is not None]
            if not pares:
                continue
            pr, rc, f1 = prec_rec(pares)
            md.append(f"| {nome} | {alvo_nome} | {pr:.3f} | {rc:.3f} | {f1:.3f} | {len(pares)} |")
            metricas.append({"metrica": "desempenho", "comparacao": nome, "dimensao": alvo_nome,
                             "valor": round(f1, 4) if not math.isnan(f1) else "", "precisao": round(pr, 4) if not math.isnan(pr) else "",
                             "recuperacao": round(rc, 4) if not math.isnan(rc) else "", "n": len(pares)})
    md.append("")

    # 3b. v0.3.4: filtro atual por formato do vídeo (Shorts têm títulos curtos e cheios de hashtags;
    # o filtro tende a errar de outro jeito neles). Só reportado: o desenho amostral não estratifica
    # por formato, então os ICs por formato serão largos se houver poucos pares.
    formatos = sorted({chave[p].get("formato") or "sem_formato" for p in gab})
    if formatos != ["sem_formato"]:
        md += ["## Filtro atual por formato do vídeo (ponderado)", "",
               "| Formato | Alvo | Precisão | Recuperação | F1 | n | relevantes no gabarito |", "|---|---|---|---|---|---|---|"]
        for fmt in formatos:
            ids = [p for p in gab if (chave[p].get("formato") or "sem_formato") == fmt]
            for i, alvo_nome in enumerate(("municipal_relevante", "regional_relevante")):
                def filtro_fmt(p, i=i):
                    k = chave[p]
                    aceito = k.get("aceito_padrao") == "True"
                    return int(aceito and k.get("municipal_evidence") not in ("", "nenhuma")) if i == 0 else int(aceito)
                pares = [(filtro_fmt(p), alvo(gab[p])[i], float(chave[p]["peso"] or 1)) for p in ids]
                pares = [x for x in pares if x[1] is not None]
                if not pares:
                    continue
                pr, rc, f1 = prec_rec(pares)
                md.append(f"| {fmt} | {alvo_nome} | {pr:.3f} | {rc:.3f} | {f1:.3f} | {len(pares)} | {sum(t for _, t, _ in pares)} |")
                metricas.append({"metrica": "desempenho_formato", "comparacao": f"filtro atual [{fmt}]", "dimensao": alvo_nome,
                                 "valor": round(f1, 4) if not math.isnan(f1) else "", "precisao": round(pr, 4) if not math.isnan(pr) else "",
                                 "recuperacao": round(rc, 4) if not math.isnan(rc) else "", "n": len(pares)})
        md.append("")

    # 4. erros do filtro por estrato (onde o filtro falha)
    md += ["## Gabarito por estrato (proporção de pares relevantes)", "", "| Estrato | n | municipal | regional |", "|---|---|---|---|"]
    por = defaultdict(list)
    for p, r in gab.items():
        por[chave[p]["estrato"]].append(alvo(r))
    for e in sorted(por):
        v = por[e]
        md.append(f"| {e} | {len(v)} | {sum(a for a, _ in v) / len(v):.0%} | {sum(b for _, b in v) / len(v):.0%} |")
    md.append("")
    md.append("Leitura: em E1 a proporção municipal é a precisão do filtro atual nesse estrato; em E3–E6 "
              "ela é a taxa de falsos negativos (E6 é o teste de elusão).")

    (d / "relatorio_calibracao.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    with (d / "metricas.csv").open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["metrica", "comparacao", "dimensao", "valor", "ic_inf", "ic_sup",
                                           "precisao", "recuperacao", "n"])
        w.writeheader(); w.writerows(metricas)
    print("\n".join(md))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
